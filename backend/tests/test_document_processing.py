from io import BytesIO

import pymupdf
import pytest
from docx import Document

from app.documents.errors import DocumentProcessingError
from app.documents.models import DocumentProcessingStatus
from app.documents.service import DocumentProcessingService
from app.main import app


@pytest.fixture
def service() -> DocumentProcessingService:
    return DocumentProcessingService(max_file_size_bytes=1_000_000)


def make_pdf() -> bytes:
    document = pymupdf.open()
    first_page = document.new_page()
    first_page.insert_text((72, 72), "First page text")
    second_page = document.new_page()
    second_page.insert_text((72, 72), "Second page text")
    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("Opening paragraph")
    document.add_paragraph("Second paragraph", style="Title")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table paragraph"
    content = BytesIO()
    document.save(content)
    return content.getvalue()


def test_extract_pdf_preserves_page_text_and_numbers(service: DocumentProcessingService) -> None:
    result = service.process("report.pdf", "application/pdf", make_pdf())

    assert result.status is DocumentProcessingStatus.READY
    assert [(segment.page_number, segment.text.strip()) for segment in result.segments] == [
        (1, "First page text"),
        (2, "Second page text"),
    ]


def test_extract_docx_preserves_paragraph_order_and_metadata(
    service: DocumentProcessingService,
) -> None:
    result = service.process(
        "report.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        make_docx(),
    )

    assert [segment.text for segment in result.segments] == [
        "Opening paragraph",
        "Second paragraph",
        "Table paragraph",
    ]
    assert [segment.paragraph_index for segment in result.segments] == [1, 2, 3]
    assert result.segments[1].paragraph_style == "Title"
    assert result.segments[2].table_index == 1
    assert result.segments[2].table_row_index == 1
    assert result.segments[2].table_cell_index == 1


def test_extract_txt_preserves_line_order(service: DocumentProcessingService) -> None:
    result = service.process("notes.txt", "text/plain; charset=utf-8", b"first\nsecond\nthird")

    assert [(segment.line_number, segment.text) for segment in result.segments] == [
        (1, "first"),
        (2, "second"),
        (3, "third"),
    ]


def test_empty_file_is_rejected(service: DocumentProcessingService) -> None:
    with pytest.raises(DocumentProcessingError) as error:
        service.process("empty.txt", "text/plain", b"")

    assert error.value.code == "empty_file"
    assert error.value.status_code == 422
    assert error.value.context == {"status": "failed"}


def test_unsupported_file_is_rejected(service: DocumentProcessingService) -> None:
    with pytest.raises(DocumentProcessingError) as error:
        service.process("image.png", "image/png", b"not empty")

    assert error.value.code == "unsupported_file_type"
    assert error.value.status_code == 415


def test_mismatched_content_type_is_rejected(service: DocumentProcessingService) -> None:
    with pytest.raises(DocumentProcessingError) as error:
        service.process("notes.txt", "application/pdf", b"plain text")

    assert error.value.code == "content_type_mismatch"
    assert error.value.status_code == 415


def test_file_over_configured_size_is_rejected() -> None:
    service = DocumentProcessingService(max_file_size_bytes=4)

    with pytest.raises(DocumentProcessingError) as error:
        service.process("notes.txt", "text/plain", b"12345")

    assert error.value.code == "file_too_large"
    assert error.value.status_code == 413


def test_extraction_failure_is_reported_structurally(
    service: DocumentProcessingService,
) -> None:
    with pytest.raises(DocumentProcessingError) as error:
        service.process("broken.pdf", "application/pdf", b"not a valid PDF")

    assert error.value.code == "extraction_failed"
    assert error.value.status_code == 422
    assert error.value.context == {"status": "failed"}
